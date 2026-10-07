# -*- coding: utf-8 -*-
"""
修复 Word 原件中"显示公式的编号未右对齐"的问题。

做法：把显示公式段落改为
    [居中制表位] <公式> [右制表位] （编号）
即在段落上设置两个制表位：
    - 居中制表位 @ 半个版心宽  -> 公式居中
    - 右制表位   @ 整个版心宽  -> 编号贴右边

背景：公式是 OMML 对象（m:oMathPara），它必须位于段落两个 run 之间；
python-docx 的 paragraph.add_run() 只能追加到末尾，因此这里直接操作 XML，
用 addprevious() 把第一个制表符插到公式前面。

用法：
    python fix_eq_number.py 0202
    python fix_eq_number.py 0202 --dry-run
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu

DIP = Path(r"D:\workspace\Jupyter_WorkSpace\DIP")
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
M_OMATHPARA = f"{{{M}}}oMathPara"


def make_tab_run():
    """构造一个只含制表符的 w:r。"""
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:tab"))
    return r


def is_tab_run(el):
    return (el.tag == qn("w:r") and el.find(qn("w:tab")) is not None
            and not (el.find(qn("w:t")) is not None
                     and (el.find(qn("w:t")).text or "").strip()))


def fix_paragraph(para, avail_cm):
    """修复单个公式段落，返回是否改动。"""
    p = para._p
    omp = p.find(M_OMATHPARA)
    if omp is None:
        return False

    # 1) 取消居中（居中会让整体居中，编号无法到右边）
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # 2) 设置两个制表位
    pf = para.paragraph_format
    for ts in list(pf.tab_stops):
        pass                      # 保留原有，仅补充
    pf.tab_stops.add_tab_stop(Cm(avail_cm / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(avail_cm), WD_TAB_ALIGNMENT.RIGHT)

    # 3) 拆出"编号 run"（公式之后、非制表符的 run）
    children = list(p)
    idx = children.index(omp)
    after = [c for c in children[idx + 1:]
             if c.tag == qn("w:r") and not is_tab_run(c)]
    if not after:
        return None               # 没有编号，交由调用方报告

    # 编号 run：保留第一个，移除其余
    numrun = after[0]
    for extra in after[1:]:
        p.remove(extra)

    # 4) 去掉公式后面原有的制表符/空白 run，重新插入
    for c in children[idx + 1:]:
        if c is numrun:
            break
        if is_tab_run(c) or (c.tag == qn("w:r") and c.find(qn("w:t")) is not None):
            t = c.find(qn("w:t"))
            if t is None or not (t.text or "").strip():
                p.remove(c)

    # 5) 清理编号文本开头的全角空格
    for t in numrun.findall(qn("w:t")):
        if t.text:
            t.text = t.text.lstrip("\u3000 ")
            t.set(qn("xml:space"), "preserve")

    # 6) 插入两个制表符
    omp.addprevious(make_tab_run())      # 公式前的居中制表符
    numrun.addprevious(make_tab_run())   # 编号前的右制表符
    return True


def process(chap, dry_run=False):
    folder = DIP / chap
    docs = [p for p in folder.glob("*.docx")
            if not p.name.startswith("~$") and "副本" not in p.name]
    if not docs:
        print(f"{chap}: 未找到原件")
        return
    path = docs[0]
    print("=" * 74)
    print(f"{chap}  {path.name}")
    print("=" * 74)

    doc = Document(str(path))
    sec = doc.sections[0]
    avail = Emu(sec.page_width - sec.left_margin - sec.right_margin).cm
    print(f"  版心宽度: {avail:.2f} cm  -> 制表位 {avail/2:.2f} / {avail:.2f}")

    n_ok = n_nonum = 0
    for para in doc.paragraphs:
        if para._p.find(M_OMATHPARA) is None:
            continue
        r = fix_paragraph(para, avail)
        if r is True:
            n_ok += 1
        elif r is None:
            n_nonum += 1

    print(f"  已修复公式段落: {n_ok}   无编号的公式段落: {n_nonum}")
    if n_nonum:
        print("  !! 有公式段落没有编号，需人工检查")

    # 校验：修复后的制表位与文本
    print("  修复后抽查：")
    shown = 0
    for para in doc.paragraphs:
        if para._p.find(M_OMATHPARA) is None:
            continue
        shown += 1
        if shown > 3:
            break
        ts = [(round(t.position.cm, 2), str(t.alignment)) for t in para.paragraph_format.tab_stops]
        txt = "".join(r.text or "" for r in para.runs)
        print(f"    制表位={ts}  编号文本={txt!r}")

    if dry_run:
        print("  [dry-run] 未写盘")
        return
    doc.save(str(path))
    print(f"  已写回: {path.name}  ({path.stat().st_size/1024:.1f} KB)")


if __name__ == "__main__":
    chap = sys.argv[1] if len(sys.argv) > 1 else "0202"
    dry = "--dry-run" in sys.argv
    process(chap, dry)
